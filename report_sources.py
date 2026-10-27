"""Bounded file/text extraction and public HTTPS source retrieval."""
from io import BytesIO
from html.parser import HTMLParser
import ipaddress
import socket
from urllib.parse import urlparse, urljoin
from zipfile import ZipFile

from docx import Document
import fitz
import requests

MAX_FILE = 10 * 1024 * 1024
MAX_TEXT = 90000


def extract_file(name: str, content: bytes):
    if len(content) > MAX_FILE:
        raise ValueError('الحد الأقصى لكل ملف هو 10 ميغابايت.')
    suffix = name.rsplit('.', 1)[-1].lower()
    try:
        if suffix == 'pdf':
            with fitz.open(stream=content, filetype='pdf') as pdf:
                if pdf.needs_pass or len(pdf) > 400:
                    raise ValueError('ملف PDF مشفّر أو يتجاوز 400 صفحة.')
                parts, size = [], 0
                for n, page in enumerate(pdf):
                    text = page.get_text()
                    size += len(text)
                    if size > MAX_TEXT:
                        raise ValueError('النص المستخرج طويل جداً؛ ارفع ملفاً أصغر.')
                    parts.append(f'[صفحة {n + 1}]\n{text}')
                result = '\n'.join(parts)
                if not any(page.strip() for page in [p.split(']\n', 1)[-1] for p in parts]):
                    raise ValueError('ملف PDF مصوّر بلا نص قابل للاستخراج؛ استخدم OCR أولاً.')
        elif suffix == 'docx':
            with ZipFile(BytesIO(content)) as archive:
                if len(archive.infolist()) > 2000 or sum(f.file_size for f in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('ملف Word أكبر من حدود المعالجة بعد فك الضغط.')
            document = Document(BytesIO(content))
            result = '\n'.join([p.text for p in document.paragraphs] + [' | '.join(c.text for c in row.cells) for table in document.tables for row in table.rows])
        elif suffix in ('txt', 'md'):
            result = content.decode('utf-8-sig')
        else:
            raise ValueError('صيغة غير مدعومة. استخدم PDF أو DOCX أو TXT أو MD.')
    except (ValueError, UnicodeError) as exc:
        raise ValueError(str(exc)) from None
    except Exception:
        # Parsers have different exception classes; never expose uploaded contents.
        raise ValueError('تعذّرت قراءة الملف. تحقق من سلامته وصيغته.') from None
    if not result.strip() or len(result) > MAX_TEXT:
        raise ValueError('الملف فارغ أو يتجاوز 90 ألف حرف مستخرج.')
    return result


class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'noscript'):
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'noscript'):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden and data.strip():
            self.parts.append(data.strip())


def validate_public_url(url):
    parsed = urlparse(url)
    if parsed.scheme != 'https' or not parsed.hostname or parsed.username or parsed.password or parsed.port not in (None, 443):
        raise ValueError('استخدم رابط HTTPS عاماً دون بيانات دخول أو منفذ مخصص.')
    try:
        # Reject private IP literals even if the machine has no DNS service.
        literal = ipaddress.ip_address(parsed.hostname)
    except ValueError:
        literal = None
    if literal is not None:
        addresses = [literal]
    else:
        try:
            addresses = [ipaddress.ip_address(a[4][0]) for a in socket.getaddrinfo(parsed.hostname, 443, type=socket.SOCK_STREAM)]
        except OSError:
            # Proxy-only cloud machines may not offer local DNS. Query a named,
            # TLS-verified public resolver; never bypass the public-IP check.
            addresses = []
            try:
                for record_type, type_id in [('A', 1), ('AAAA', 28)]:
                    response = requests.get('https://dns.google/resolve',
                                            params={'name': parsed.hostname, 'type': record_type}, timeout=(10, 10))
                    response.raise_for_status()
                    data = response.json()
                    if data.get('Status') != 0:
                        break
                    addresses.extend(ipaddress.ip_address(a['data']) for a in data.get('Answer', []) if a.get('type') == type_id)
                if not addresses:
                    raise ValueError('No public address records')
            except (requests.RequestException, ValueError, KeyError, TypeError, AttributeError):
                raise ValueError('تعذّر التحقق من عنوان المصدر: DNS المحلي غير متاح. اسمح بنطاق dns.google في إعدادات الشبكة، أو أرفق نص المصدر مباشرة.') from None
    if not addresses or any(not address.is_global for address in addresses):
        raise ValueError('الروابط المحلية وعناوين الشبكات الخاصة غير مسموحة.')
    return url


def fetch_source(url):
    try:
        for _ in range(4):
            validate_public_url(url)
            with requests.get(url, timeout=(10, 20), stream=True, allow_redirects=False,
                              headers={'User-Agent': 'RiwaqReport/1.0'}) as response:
                if response.is_redirect:
                    url = urljoin(url, response.headers.get('Location', ''))
                    continue
                response.raise_for_status()
                body = bytearray()
                for block in response.iter_content(32768):
                    body.extend(block)
                    if len(body) > 3 * 1024 * 1024:
                        raise ValueError('المصدر الإلكتروني يتجاوز حد 3 ميغابايت.')
                kind = response.headers.get('Content-Type', '').lower()
                if 'application/pdf' in kind:
                    return extract_file('source.pdf', bytes(body))
                if 'html' not in kind and not kind.startswith('text/'):
                    raise ValueError('الرابط لا يعيد صفحة نصية أو ملف PDF مدعوماً.')
                text = bytes(body).decode(response.encoding if response.encoding and response.encoding.lower() != 'iso-8859-1' else 'utf-8', errors='replace')
                if 'html' in kind:
                    parser = TextParser()
                    parser.feed(text)
                    text = '\n'.join(parser.parts)
                if not text.strip() or len(text) > MAX_TEXT:
                    raise ValueError('المصدر فارغ أو طويل جداً. أرفق الجزء المطلوب كنص بدلاً منه.')
                return text
        raise ValueError('تجاوز المصدر حد إعادة التوجيه.')
    except requests.RequestException:
        raise ValueError('تعذّر جلب المصدر؛ تحقق من إتاحته وسياسة الشبكة. يمكنك إرفاق نصه مباشرة.') from None


class SearchParser(HTMLParser):
    """Read the public HTML search page without a commercial search API."""
    def __init__(self):
        super().__init__()
        self.results = []
        self.current = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'a' and 'result__a' in attrs.get('class', ''):
            self.current = [attrs.get('href', ''), '']

    def handle_data(self, data):
        if self.current is not None:
            self.current[1] += data

    def handle_endtag(self, tag):
        if tag == 'a' and self.current is not None:
            self.results.append(self.current)
            self.current = None


def search_sources(query):
    from urllib.parse import parse_qs
    url = 'https://html.duckduckgo.com/html/'
    validate_public_url(url)
    try:
        with requests.get(url, params={'q': query}, timeout=(10, 20), stream=True,
                          headers={'User-Agent': 'RiwaqReport/1.0'}, allow_redirects=False) as response:
            response.raise_for_status()
            if response.is_redirect:
                raise ValueError('أعاد محرك البحث توجيهاً غير متوقع. أضف روابط المصادر يدوياً.')
            content = bytearray()
            for block in response.iter_content(32768):
                content.extend(block)
                if len(content) > 2 * 1024 * 1024:
                    raise ValueError('صفحة البحث تتجاوز حد المعالجة.')
        parser = SearchParser()
        parser.feed(bytes(content).decode('utf-8', errors='replace'))
    except requests.RequestException:
        raise ValueError('البحث المجاني غير متاح حالياً. أرفق مصادر أو روابط مباشرة بدلاً منه.') from None
    sources, seen = [], set()
    for raw_url, title in parser.results[:10]:
        target = parse_qs(urlparse(raw_url).query).get('uddg', [raw_url])[0]
        if not target.startswith('https://') or target in seen:
            continue
        seen.add(target)
        try:
            text = fetch_source(target)
        except ValueError:
            continue
        sources.append((title, target, text))
        if len(sources) == 3:
            break
    if not sources:
        raise ValueError('لم نجد مصادر عامة قابلة للقراءة. أضف ملفات أو نصوصاً أو روابط موثوقة يدوياً.')
    return sources
