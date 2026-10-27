"""Render Markdown as an editable, styled Word document with Arabic bidi support."""
from io import BytesIO
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt

from report_engine import ReportRequest, report_date


def set_rtl(paragraph, rtl):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1' if rtl else '0')
    paragraph._p.get_or_add_pPr().append(bidi)


def style_run(run, rtl):
    run.font.name = 'Arial'
    fonts = run._element.get_or_add_rPr()
    family = OxmlElement('w:rFonts')
    for name in ('ascii', 'hAnsi', 'cs'):
        family.set(qn('w:' + name), 'Arial')
    fonts.append(family)
    if rtl:
        bidi = OxmlElement('w:rtl')
        bidi.set(qn('w:val'), '1')
        fonts.append(bidi)


def inline(paragraph, children, rtl):
    bold = italic = False
    for token in children or []:
        if token.type in ('strong_open', 'strong_close'):
            bold = token.type.endswith('open')
        elif token.type in ('em_open', 'em_close'):
            italic = token.type.endswith('open')
        elif token.type in ('text', 'code_inline'):
            run = paragraph.add_run(token.content)
            run.bold, run.italic = bold, italic
            style_run(run, rtl and not bool(re.fullmatch(r'https?://\S+', token.content)))
            if token.type == 'code_inline':
                run.font.name = 'Consolas'
        elif token.type in ('softbreak', 'hardbreak'):
            paragraph.add_run().add_break()
        elif token.type == 'link_close':
            # Keep link destinations visible and usable even when Word strips fields.
            pass
        elif token.type == 'link_open':
            href = token.attrGet('href') or ''
            if href.startswith(('https://', 'http://')):
                run = paragraph.add_run(f' ({href}) ')
                style_run(run, False)
                run.font.color.rgb = RGBColor.from_string('157D73')


def shade(paragraph):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), 'EAF5F2')
    paragraph._p.get_or_add_pPr().append(shading)


def field(paragraph, instruction):
    run = paragraph.add_run()
    for tag, value in [('begin', None), (None, instruction), ('end', None)]:
        element = OxmlElement('w:instrText' if tag is None else 'w:fldChar')
        if tag is not None:
            element.set(qn('w:fldCharType'), tag)
        else:
            element.text = value
        run._r.append(element)


def build_docx(markdown: str, spec: ReportRequest, citations=None):
    if not markdown.strip():
        raise ValueError('لا يمكن تصدير تقرير فارغ.')
    if len(markdown) > 200000:
        raise ValueError('النص يتجاوز الحد الأقصى للتصدير.')
    rtl = spec.language != 'English'
    tokens = MarkdownIt('commonmark').enable('table').parse(markdown)
    document = Document()
    document.core_properties.title = spec.topic
    document.core_properties.subject = spec.kind
    document.core_properties.author = 'رِواق | Riwaq'
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    section.different_first_page_header_footer = True
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(12 if name == 'Normal' else 13)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.35
    for name, size in [('Title', 30), ('Heading 1', 21), ('Heading 2', 16), ('Heading 3', 14)]:
        document.styles[name].font.size = Pt(size)
        document.styles[name].font.color.rgb = RGBColor.from_string('167D73')
    title = document.add_paragraph(spec.topic, 'Title')
    set_rtl(title, rtl)
    title.paragraph_format.space_before = Pt(110)
    english_kinds = {'شرح تعليمي مفصل': 'Detailed educational explanation', 'تقرير تنفيذي': 'Executive report', 'دراسة أكاديمية': 'Academic study', 'ملخص شامل': 'Comprehensive summary', 'دليل إرشادي': 'Practical guide'}
    english_levels = {'مبتدئ': 'Beginner', 'متوسط': 'Intermediate', 'محترف وخبير': 'Professional / Expert'}
    subtitle = document.add_paragraph(spec.kind if rtl else english_kinds.get(spec.kind, spec.kind), 'Subtitle')
    set_rtl(subtitle, rtl)
    length_label = {'موجز: 3–5 صفحات': 'Short: 3–5 pages', 'متوسط: 6–10 صفحات': 'Medium: 6–10 pages', 'تفصيلي شامل: 10+ صفحات': 'Comprehensive: 10+ pages'}
    labels = [('المستوى المستهدف', spec.level), ('اللغة', spec.language), ('الطول المستهدف', spec.length), ('التاريخ', report_date())] if rtl else [('Audience', english_levels.get(spec.level, spec.level)), ('Language', 'English'), ('Target length', length_label.get(spec.length, spec.length)), ('Date', report_date())]
    for label, value in labels:
        p = document.add_paragraph(f'{label}: {value}')
        set_rtl(p, rtl)
    document.add_page_break()
    p = document.add_paragraph('فهرس المحتويات' if rtl else 'Table of contents', 'Heading 1')
    set_rtl(p, rtl)
    headings = [(i, int(t.tag[1]), tokens[i + 1].content) for i, t in enumerate(tokens) if t.type == 'heading_open' and int(t.tag[1]) > 1]
    if headings:
        toc = document.add_table(rows=1, cols=2)
        toc.style = 'Light Shading Accent 1'
        toc.rows[0].cells[0].text = 'القسم' if rtl else 'Section'
        toc.rows[0].cells[1].text = 'العنوان' if rtl else 'Heading'
        for number, (_, level, text) in enumerate(headings, 1):
            cells = toc.add_row().cells
            cells[0].text, cells[1].text = str(number), ('    ' * (level - 2)) + text
        for row in toc.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    set_rtl(paragraph, rtl)
        if rtl:
            toc._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
    else:
        document.add_paragraph('لا توجد عناوين فرعية في النص.' if rtl else 'No section headings in this text.')
    document.add_page_break()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('رِواق  •  ' if rtl else 'Riwaq  •  ')
    field(footer, ' PAGE ')
    header = section.header.paragraphs[0]
    header.add_run(spec.topic)
    set_rtl(header, rtl)
    lists = []
    quote = 0
    i = 0
    while i < len(tokens):
        t = tokens[i]
        if t.type == 'heading_open':
            level = int(t.tag[1])
            if level > 1:
                p = document.add_paragraph(style=f'Heading {min(level - 1, 3)}')
                set_rtl(p, rtl)
                inline(p, tokens[i + 1].children, rtl)
            i += 3
            continue
        if t.type in ('bullet_list_open', 'ordered_list_open'):
            lists.append('List Number' if t.type == 'ordered_list_open' else 'List Bullet')
        elif t.type in ('bullet_list_close', 'ordered_list_close'):
            lists.pop()
        elif t.type == 'blockquote_open':
            quote += 1
        elif t.type == 'blockquote_close':
            quote -= 1
        elif t.type == 'paragraph_open':
            p = document.add_paragraph(style=lists[-1] if lists else 'Normal')
            set_rtl(p, rtl)
            inline(p, tokens[i + 1].children, rtl)
            if quote:
                shade(p)
            i += 3
            continue
        elif t.type in ('fence', 'code_block'):
            p = document.add_paragraph(t.content)
            set_rtl(p, False)
            shade(p)
            for run in p.runs:
                run.font.name = 'Consolas'
        elif t.type == 'table_open':
            rows, cells = [], []
            while tokens[i].type != 'table_close':
                item = tokens[i]
                if item.type == 'tr_open':
                    cells = []
                elif item.type in ('th_open', 'td_open'):
                    cells.append(tokens[i + 1])
                elif item.type == 'tr_close':
                    rows.append(cells)
                i += 1
            if rows:
                table = document.add_table(rows=0, cols=max(len(r) for r in rows))
                table.style = 'Light Shading Accent 1'
                if rtl:
                    table._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
                for n, row in enumerate(rows):
                    added = table.add_row()
                    if n == 0:
                        added._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                    for cell, content in zip(added.cells, row):
                        set_rtl(cell.paragraphs[0], rtl)
                        inline(cell.paragraphs[0], content.children, rtl)
        i += 1
    if citations:
        p = document.add_paragraph('روابط المصادر المسترجعة' if rtl else 'Retrieved source links', 'Heading 1')
        set_rtl(p, rtl)
        for source in citations:
            p = document.add_paragraph(f"{source.get('title', '')}\n{source['url']}")
            set_rtl(p, rtl)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
