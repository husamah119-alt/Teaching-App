from io import BytesIO
from pathlib import Path
from unittest.mock import Mock
from zipfile import ZipFile
import socket

import pytest
from docx import Document
from streamlit.testing.v1 import AppTest

from document_export import build_docx
from report_engine import ReportRequest, SOURCE_MODES, build_prompt, demo_report, generate_report
from report_sources import extract_file, fetch_source, validate_public_url


def test_word_contains_cover_toc_tables_rtl_and_editable_content():
    data = build_docx(demo_report(), ReportRequest(topic='تقرير التعلّم'))
    doc = Document(BytesIO(data))
    assert doc.paragraphs[0].text == 'تقرير التعلّم'
    assert any(p.text == 'فهرس المحتويات' for p in doc.paragraphs)
    assert len(doc.tables) == 2
    assert any('الاسترجاع النشط' in cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert doc.tables[1].rows[1].cells[0].text == 'الاسترجاع'
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read('word/document.xml').decode()
        assert 'w:bidi' in xml and 'w:rtl' in xml and 'w:bidiVisual' in xml
        assert 'PAGE' in archive.read('word/footer1.xml').decode()
    english = Document(BytesIO(build_docx('## Introduction\nA **bold** introduction.', ReportRequest(topic='Test', language='English'))))
    assert any(p.text == 'Table of contents' for p in english.paragraphs)
    assert any(r.bold and r.text == 'bold' for p in english.paragraphs for r in p.runs)


def test_source_extraction_text_docx_pdf_and_invalid():
    assert extract_file('source.txt', 'نص عربي'.encode()) == 'نص عربي'
    doc = Document()
    doc.add_paragraph('Source paragraph')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Source cell'
    data = BytesIO()
    doc.save(data)
    assert 'Source cell' in extract_file('source.docx', data.getvalue())
    from pipeline import demo_pdf
    assert '[صفحة 1]' in extract_file('source.pdf', demo_pdf())
    for name, content in [('bad.pdf', b'bad'), ('bad.docx', b'bad'), ('empty.txt', b''), ('bad.exe', b'x')]:
        with pytest.raises(ValueError):
            extract_file(name, content)


def test_private_urls_are_not_fetched(monkeypatch):
    monkeypatch.setattr(socket, 'getaddrinfo', lambda *a, **kw: [(socket.AF_INET, socket.SOCK_STREAM, 6, '', ('127.0.0.1', 443))])
    get = Mock()
    monkeypatch.setattr('report_sources.requests.get', get)
    with pytest.raises(ValueError):
        fetch_source('https://example.com/private')
    get.assert_not_called()
    for url in ['http://example.com', 'https://user:pass@example.com', 'https://example.com:8080']:
        with pytest.raises(ValueError):
            validate_public_url(url)


def test_redirect_to_private_address_rejected(monkeypatch):
    monkeypatch.setattr(socket, 'getaddrinfo', lambda host, *a, **kw: [(socket.AF_INET, socket.SOCK_STREAM, 6, '', ('127.0.0.1' if host == 'localhost' else '93.184.216.34', 443))])
    response = Mock(is_redirect=True, headers={'Location': 'https://localhost/secret'})
    response.__enter__ = Mock(return_value=response)
    response.__exit__ = Mock(return_value=False)
    get = Mock(return_value=response)
    monkeypatch.setattr('report_sources.requests.get', get)
    with pytest.raises(ValueError):
        fetch_source('https://example.com')
    assert get.call_count == 1


def test_prompt_preserves_spec_and_source_constraints():
    spec = ReportRequest(topic='الطاقة الشمسية', sources=SOURCE_MODES[2], level='محترف وخبير', language='English')
    prompt = build_prompt(spec, 'source.pdf [صفحة 2] data')
    assert 'Use ONLY' in prompt and 'source.pdf' in prompt and 'English' in prompt
    with pytest.raises(ValueError):
        generate_report(spec, '')
    with pytest.raises(ValueError):
        ReportRequest(topic='').validate()


def test_free_source_analysis_never_uses_generation_service(monkeypatch):
    post = Mock(side_effect=AssertionError('No paid or external generation calls'))
    monkeypatch.setattr('requests.post', post)
    text, citations = generate_report(ReportRequest(topic='الطاقة الشمسية', sources=SOURCE_MODES[2]),
        'SOURCE: book.txt\nتحول الألواح الشمسية الطاقة الضوئية إلى كهرباء.\n\nتستخدم البطاريات لتخزين الطاقة.')
    assert 'تحول الألواح الشمسية' in text and 'book.txt' in text
    assert 'استخراجي' in text and citations == []
    document = Document(BytesIO(build_docx(text, ReportRequest(topic='الطاقة الشمسية'))))
    assert len(document.tables) == 2
    assert document.tables[1].rows[1].cells[0].text == 'book.txt'
    post.assert_not_called()
    with pytest.raises(ValueError, match='أرفق'):
        generate_report(ReportRequest(topic='موضوع بلا مصادر'), '')


def test_optional_local_inference_and_truncation(monkeypatch):
    model = Mock()
    model.tokenize.return_value = list(range(100))
    model.create_chat_completion.return_value = {'choices': [{'finish_reason': 'stop', 'message': {'content': '## المقدمة\nنص محلي.'}}]}
    monkeypatch.setattr('report_engine.load_model', lambda: model)
    text, _ = generate_report(ReportRequest(topic='الماء'), '', use_model=True)
    assert 'نص محلي' in text
    assert 'الماء' in model.create_chat_completion.call_args.kwargs['messages'][0]['content']
    model.tokenize.return_value = list(range(8000))
    with pytest.raises(ValueError, match='طويلة'):
        generate_report(ReportRequest(topic='الماء'), 'source', use_model=True)
    model.tokenize.return_value = [1]
    model.create_chat_completion.return_value = {'choices': [{'finish_reason': 'length'}]}
    with pytest.raises(ValueError, match='لم يكتمل'):
        generate_report(ReportRequest(topic='الماء'), '', use_model=True)


def test_report_ui_demo_edit_export_and_clear():
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'), default_timeout=30).run()
    assert not app.exception
    next(b for b in app.button if b.label == 'فتح نموذج تنسيق').click().run()
    assert not app.exception
    assert len(app.get('download_button')) == 2
    editor = next(t for t in app.text_area if t.label == 'محتوى التقرير بصيغة Markdown')
    editor.input('## تقرير معدّل\nتم تعديل النص بنجاح.').run()
    assert 'تم تعديل النص' in app.session_state['report_text']
    next(b for b in app.button if b.label == 'مسح التقرير من الجلسة').click().run()
    assert not app.exception
    assert len(app.get('download_button')) == 0


def test_report_ui_has_no_credentials_and_requests_sources():
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'), default_timeout=30).run()
    assert not any('مفتاح' in t.label or 'OpenAI' in t.label for t in app.text_input)
    next(t for t in app.text_input if t.label.startswith('موضوع البحث')).input('اختبار')
    next(b for b in app.button if b.label.startswith('إنشاء الشرح')).click().run()
    assert not app.exception
    assert any('أضف ملفاً' in error.value for error in app.error)


def test_report_ui_generation_with_supplied_sources(monkeypatch):
    seen = {}

    def provider(spec, evidence, *, use_model=False):
        seen.update(spec=spec, evidence=evidence)
        return '## المقدمة\nتقرير يستند إلى النص المرفق.\n\n## المصادر والمراجع\nنص أرفقه المستخدم.', []

    monkeypatch.setattr('report_engine.generate_report', provider)
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'), default_timeout=30).run()
    next(t for t in app.text_input if t.label.startswith('موضوع البحث')).input('التعلّم النشط')
    next(s for s in app.selectbox if s.label == 'المصادر المعتمدة').select(SOURCE_MODES[2])
    next(t for t in app.text_area if t.label.startswith('نصوص المصادر')).input('المصدر: نشاط تعليمي تجريبي.')
    next(b for b in app.button if b.label.startswith('إنشاء الشرح')).click().run()
    assert not app.exception
    assert seen['spec'].topic == 'التعلّم النشط'
    assert 'نشاط تعليمي تجريبي' in seen['evidence']
    assert len(app.get('download_button')) == 2
    assert app.session_state['is_demo'] is False
    assert app.session_state['report_manifest'][0]['المصدر'] == 'نص أرفقه المستخدم'


def test_bounded_html_retrieval_and_extraction(monkeypatch):
    monkeypatch.setattr(socket, 'getaddrinfo', lambda *a, **kw: [(socket.AF_INET, socket.SOCK_STREAM, 6, '', ('93.184.216.34', 443))])
    response = Mock(is_redirect=False, headers={'Content-Type': 'text/html'}, encoding='utf-8')
    response.iter_content.return_value = [b'<h1>Source title</h1><script>hidden text</script><p>Evidence.</p>']
    response.__enter__ = Mock(return_value=response)
    response.__exit__ = Mock(return_value=False)
    monkeypatch.setattr('report_sources.requests.get', Mock(return_value=response))
    assert fetch_source('https://example.com') == 'Source title\nEvidence.'
    response.iter_content.return_value = [b'x' * (3 * 1024 * 1024 + 1)]
    with pytest.raises(ValueError, match='3'):
        fetch_source('https://example.com')
    with pytest.raises(ValueError, match='10'):
        extract_file('big.txt', b'x' * (10 * 1024 * 1024 + 1))


def test_proxy_only_dns_fallback_keeps_public_address_checks(monkeypatch):
    def no_dns(*args, **kwargs):
        raise socket.gaierror('unavailable')
    monkeypatch.setattr(socket, 'getaddrinfo', no_dns)
    response = Mock(status_code=200, json=lambda: {'Status': 0, 'Answer': [{'type': 1, 'data': '93.184.216.34'}]})
    get = Mock(return_value=response)
    monkeypatch.setattr('report_sources.requests.get', get)
    assert validate_public_url('https://example.com') == 'https://example.com'
    assert get.call_args.args[0] == 'https://dns.google/resolve'
    response.json = lambda: {'Status': 0, 'Answer': [{'type': 1, 'data': '10.0.0.1'}]}
    with pytest.raises(ValueError, match='الخاصة'):
        validate_public_url('https://example.com')
    get.reset_mock()
    with pytest.raises(ValueError):
        validate_public_url('https://127.0.0.1')
    get.assert_not_called()


def test_free_report_can_be_created_and_exported_without_model_or_network(monkeypatch):
    monkeypatch.setattr('requests.post', Mock(side_effect=AssertionError('No remote generation')))
    app = AppTest.from_file(str(Path(__file__).parents[1] / 'app.py'), default_timeout=30).run()
    next(t for t in app.text_input if t.label.startswith('موضوع البحث')).input('التعلّم')
    next(t for t in app.text_area if t.label.startswith('نصوص المصادر')).input('يساعد التعلم النشط الطالب على تطبيق المفاهيم ومراجعة فهمه للنص. هذه مادة مرفقة للاختبار.')
    next(b for b in app.button if b.label.startswith('إنشاء الشرح')).click().run()
    assert not app.exception and not app.error
    assert 'يساعد التعلم النشط' in app.session_state['report_text']
    assert 'تحليل استخراجي' in app.session_state['report_text']
    assert len(app.get('download_button')) == 2


def test_short_source_is_preserved_alongside_long_source():
    text, _ = generate_report(ReportRequest(topic='موضوع'), 'SOURCE: short.txt\nقصير\nSOURCE: long.txt\nهذا مصدر ثانٍ يحتوي على فقرة طويلة بما يكفي للاختيار والقراءة.')
    assert 'قصير' in text and 'short.txt' in text


@pytest.mark.parametrize('checksum_field', ['oid', 'sha256'])
def test_official_model_installer_verifies_artifacts_and_preserves_existing(tmp_path, monkeypatch, checksum_field):
    import hashlib
    from scripts import setup_local_model as installer
    payload = b'GGUFverified test artifact'
    digest = hashlib.sha256(payload).hexdigest()
    monkeypatch.setattr(installer, 'ROOT', tmp_path)
    metadata = {'sha': 'a' * 40, 'siblings': [{'rfilename': installer.FILENAME, 'lfs': {checksum_field: digest}}]}
    first = Mock(json=lambda: metadata)
    second = Mock()
    second.iter_content.return_value = [payload]
    second.__enter__ = Mock(return_value=second)
    second.__exit__ = Mock(return_value=False)
    get = Mock(side_effect=[first, second])
    monkeypatch.setattr(installer.requests, 'get', get)
    installer.install()
    path = tmp_path / 'models' / installer.FILENAME
    assert path.read_bytes() == payload
    path.write_bytes(b'existing user file')
    get.side_effect = [first]
    with pytest.raises(ValueError, match='Existing model differs'):
        installer.install()
    assert path.read_bytes() == b'existing user file'
